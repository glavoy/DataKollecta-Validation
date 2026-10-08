#!/usr/bin/env python3
"""Build the controlled Word documents in docs/ from the Markdown sources in src/.

Each source starts with a YAML front-matter block:

    ---
    doc_id: DK-VAL-001
    title: Validation Plan
    version: "1.0"
    orientation: portrait        # or landscape
    approvers:                   # optional; defaults to Author / Reviewer / Approver
      - [Author, System Owner / Developer]
      - [Reviewer, Quality Assurance]
      - [Approver, Head of Data Management]
    ---

The build adds, in this order: a cover page with the document-control block and
the signature table, the revision history, then the body. Every page carries the
document ID, title, version and "Page X of Y".

Requires pandoc (>= 2.19) and python-docx.

    python3 tools/build_docs.py            # all documents
    python3 tools/build_docs.py DK-VAL-002 # one document
"""
from __future__ import annotations

import copy
import re
import subprocess
import sys
import tempfile
from pathlib import Path

import yaml
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
SRC = ROOT / "src"
OUT = ROOT / "docs"
REFERENCE = ROOT / "tools" / "reference.docx"

ORGANISATION = "IDRC"
SYSTEM = "DataKollecta"
ACCENT = RGBColor(0x1F, 0x3A, 0x5F)

DEFAULT_APPROVERS = [
    ["Author", "System Owner / Lead Developer"],
    ["Reviewer", "Quality Assurance (independent of development)"],
    ["Approver", "Head of Data Management"],
]

PAGE_BREAK = '\n```{=openxml}\n<w:p><w:r><w:br w:type="page"/></w:r></w:p>\n```\n'


# --------------------------------------------------------------------------- #
# Reference document (styles)
# --------------------------------------------------------------------------- #
def make_reference() -> None:
    raw = subprocess.run(
        ["pandoc", "--print-default-data-file", "reference.docx"],
        check=True, capture_output=True).stdout
    REFERENCE.write_bytes(raw)
    doc = Document(str(REFERENCE))
    styles = doc.styles

    def font(style_name: str, size: float, bold: bool | None = None,
             color: RGBColor | None = None, name: str = "Calibri") -> None:
        try:
            st = styles[style_name]
        except KeyError:
            return
        st.font.name = name
        rpr = st.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.append(rfonts)
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), name)
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
            if rfonts.get(qn(attr)) is not None:
                del rfonts.attrib[qn(attr)]
        st.font.size = Pt(size)
        if bold is not None:
            st.font.bold = bold
        if color is not None:
            st.font.color.rgb = color

    for body in ("Normal", "Body Text", "First Paragraph", "Compact"):
        font(body, 10.5)
    font("Compact", 9)
    font("Verbatim Char", 9, name="Consolas")
    font("Source Code", 8.5, name="Consolas")
    font("Title", 22, True, ACCENT)
    font("Subtitle", 14, False, ACCENT)
    font("Heading 1", 15, True, ACCENT)
    font("Heading 2", 12.5, True, ACCENT)
    font("Heading 3", 11, True, ACCENT)
    font("Heading 4", 10.5, True, ACCENT)
    for body in ("Body Text", "First Paragraph"):
        try:
            styles[body].paragraph_format.space_after = Pt(6)
            styles[body].paragraph_format.space_before = Pt(0)
        except KeyError:
            pass
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
    doc.save(str(REFERENCE))


# --------------------------------------------------------------------------- #
# Source -> Markdown with cover
# --------------------------------------------------------------------------- #
def read_source(path: Path) -> tuple[dict, str]:
    text = path.read_text(encoding="utf-8")
    m = re.match(r"^---\n(.*?)\n---\n", text, re.S)
    if not m:
        raise SystemExit(f"{path}: missing front matter")
    meta = yaml.safe_load(m.group(1))
    for key in ("doc_id", "title", "version"):
        if key not in meta:
            raise SystemExit(f"{path}: front matter lacks {key}")
    return meta, text[m.end():]


def cover_markdown(meta: dict) -> str:
    approvers = meta.get("approvers") or DEFAULT_APPROVERS
    rows = "\n".join(
        f"| **{a}** | {r} | | | |" for a, r in approvers)
    history = meta.get("history") or [
        [meta["version"], "", "First issue. Retrospective validation baseline.", ""]]
    hist_rows = "\n".join(f"| {v} | {d} | {c} | {a} |" for v, d, c, a in history)
    return f"""
::: {{custom-style="Title"}}
{SYSTEM}
:::

::: {{custom-style="Subtitle"}}
{meta['title']}
:::

| Document control | |
|:--------------------|:-------------------------------------------------------|
| **Document ID** | {meta['doc_id']} |
| **Version** | {meta['version']} |
| **Status** | {meta.get('status', 'Approved for use once signed below')} |
| **Organisation** | {ORGANISATION} |
| **System** | {SYSTEM} electronic data capture system (field application, web portal and server, survey tools) |
| **Effective date** | Date of final approval signature |
| **Supersedes** | {meta.get('supersedes', 'None (first issue)')} |

## Approval

By signing, each person confirms that they have reviewed this document and that it is
accurate and fit for its stated purpose. The document is effective from the date of the
last signature.

| Signatory | Role | Name | Signature | Date |
|:-----------|:-----------------------------|:-------------------|:-------------------|:-----------|
{rows}

## Revision history

| Version | Date | Description of change | Author |
|:--------|:---------|:-----------------------------------------------|:-----------|
{hist_rows}
{PAGE_BREAK}
"""


# --------------------------------------------------------------------------- #
# Post-processing
# --------------------------------------------------------------------------- #
def _field(paragraph, instr: str) -> None:
    run = paragraph.add_run()
    for tag, text in (("begin", None), ("instr", instr), ("separate", None),
                      ("text", "1"), ("end", None)):
        if tag == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
            run._r.append(el)
        elif tag == "text":
            t = OxmlElement("w:t")
            t.text = text
            run._r.append(t)
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
            run._r.append(el)


def _borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    existing = tbl_pr.find(qn("w:tblBorders"))
    if existing is not None:
        tbl_pr.remove(existing)
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "8C8C8C")
        borders.append(el)
    tbl_pr.append(borders)
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    # Header row: shade and repeat on every page.
    if table.rows:
        first = table.rows[0]
        tr_pr = first._tr.get_or_add_trPr()
        rep = OxmlElement("w:tblHeader")
        tr_pr.append(rep)
        for cell in first.cells:
            tc_pr = cell._tc.get_or_add_tcPr()
            old = tc_pr.find(qn("w:shd"))
            if old is not None:
                tc_pr.remove(old)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:val"), "clear")
            shd.set(qn("w:color"), "auto")
            shd.set(qn("w:fill"), "DCE4EE")
            tc_pr.append(shd)
    # Rows never split across pages, so a signature cell stays whole.
    for row in table.rows:
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)


# Schema order of the property elements this script touches. Word (unlike
# LibreOffice) rejects a document whose property children are out of order.
_ORDER = {
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize",
              "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd",
              "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook",
              "tblCaption", "tblDescription"],
    "trPr": ["cnfStyle", "divId", "gridBefore", "gridAfter", "wBefore", "wAfter",
             "cantSplit", "trHeight", "tblHeader", "tblCellSpacing", "jc", "hidden"],
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd",
             "noWrap", "tcMar", "textDirection", "tcFitText", "vAlign", "hideMark"],
    "pPr": ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
            "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd", "tabs",
            "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
            "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
            "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
            "suppressOverlap", "jc", "textDirection", "textAlignment",
            "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr", "sectPr",
            "pPrChange"],
}


def _sort_props(root) -> None:
    for name, order in _ORDER.items():
        rank = {qn(f"w:{t}"): i for i, t in enumerate(order)}
        for el in root.iter(qn(f"w:{name}")):
            children = list(el)
            seen, unique = set(), []
            for c in children:  # drop duplicate singletons (e.g. two tblLayout)
                if c.tag in seen and c.tag in rank:
                    continue
                seen.add(c.tag)
                unique.append(c)
            for c in children:
                el.remove(c)
            for c in sorted(unique, key=lambda c: rank.get(c.tag, len(order))):
                el.append(c)


def _full_width(table, text_width_emu: int) -> None:
    """Stretch every table to the text width.

    Pandoc only writes relative column widths for tables whose source lines are
    wider than --columns; other tables come out auto-fitted and narrow. Widths
    are taken from pandoc's grid where it wrote one, otherwise from the longest
    cell text in each column, and then scaled to fill the page.
    """
    tbl = table._tbl
    grid = tbl.tblGrid
    cols = grid.findall(qn("w:gridCol")) if grid is not None else []
    ncols = len(table.columns)
    weights = []
    for i in range(ncols):
        w = 0
        if i < len(cols) and cols[i].get(qn("w:w")):
            try:
                w = int(cols[i].get(qn("w:w")))
            except ValueError:
                w = 0
        weights.append(w)
    if not all(weights):
        weights = []
        for i in range(ncols):
            longest = max((len(row.cells[i].text) for row in table.rows), default=1)
            weights.append(min(max(longest, 6), 60))
    total = sum(weights) or 1
    twips_total = int(text_width_emu / 635)  # 1 twip = 635 EMU
    widths = [max(int(twips_total * w / total), 400) for w in weights]
    for i, col in enumerate(cols):
        col.set(qn("w:w"), str(widths[i]))
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i >= len(widths):
                continue
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.insert(0, tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[i]))


def _cover_break(doc) -> None:
    """Replace the cover's page-break paragraph with page-break-before on the
    next paragraph, so a cover that exactly fills its page leaves no blank page."""
    body = doc.element.body
    for p in body.iter(qn("w:p")):
        brs = p.findall(".//" + qn("w:br"))
        if any(b.get(qn("w:type")) == "page" for b in brs) and not p.findall(".//" + qn("w:t")):
            nxt = p.getnext()
            while nxt is not None and nxt.tag != qn("w:p"):
                nxt = nxt.getnext()
            if nxt is not None:
                ppr = nxt.find(qn("w:pPr"))
                if ppr is None:
                    ppr = OxmlElement("w:pPr")
                    nxt.insert(0, ppr)
                pbb = OxmlElement("w:pageBreakBefore")
                ppr.append(pbb)
                p.getparent().remove(p)
            return


def _keep_together(doc) -> None:
    """Keep a test case's header table, its steps' first row and its result table
    on the same page as their neighbours, so no table is stranded at a page foot."""
    def keep(row):
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.keep_with_next = True

    for table in doc.tables:
        if not table.rows:
            continue
        first = table.rows[0].cells[0].text.strip()
        header = [c.text.strip().lower() for c in table.rows[0].cells]
        if first == "Verifies":
            for row in table.rows:
                keep(row)
        elif "actual result" in header:
            keep(table.rows[0])
            if len(table.rows) > 1:
                keep(table.rows[1])
            keep(table.rows[-1])
        elif first in ("Signatory", "Role") and "signature" in header:
            for row in table.rows[:-1]:
                keep(row)


def _tall_signature_rows(doc) -> None:
    """Give signature/date/result cells room for a handwritten entry."""
    for table in doc.tables:
        if not table.rows:
            continue
        header = [c.text.strip().lower() for c in table.rows[0].cells]
        if "signature" in header or "actual result" in header or "pass/fail" in header:
            for row in table.rows[1:]:
                row.height = Cm(1.0)


def postprocess(path: Path, meta: dict) -> None:
    doc = Document(str(path))
    landscape = meta.get("orientation") == "landscape"
    for sec in doc.sections:
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Cm(29.7), Cm(21.0)
        else:
            sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
        sec.top_margin = sec.bottom_margin = Cm(2.0)
        sec.left_margin = sec.right_margin = Cm(2.0)
        sec.header_distance = sec.footer_distance = Cm(1.0)

        header = sec.header
        hp = header.paragraphs[0] if header.paragraphs else header.add_paragraph()
        hp.text = ""
        r = hp.add_run(f"{SYSTEM} — {meta['title']}")
        r.font.size = Pt(8.5)
        r.font.color.rgb = ACCENT
        hp.add_run("\t")
        r = hp.add_run(f"{meta['doc_id']}  ·  Version {meta['version']}")
        r.font.size = Pt(8.5)
        r.font.color.rgb = ACCENT
        width = sec.page_width - sec.left_margin - sec.right_margin
        hp.paragraph_format.tab_stops.add_tab_stop(width, alignment=2)  # right

        footer = sec.footer
        fp = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
        fp.text = ""
        r = fp.add_run(f"{ORGANISATION} — controlled document. Printed copies are uncontrolled "
                       f"unless signed.\t")
        r.font.size = Pt(8)
        fp.paragraph_format.tab_stops.add_tab_stop(width, alignment=2)
        r = fp.add_run("Page ")
        r.font.size = Pt(8)
        _field(fp, "PAGE")
        r = fp.add_run(" of ")
        r.font.size = Pt(8)
        _field(fp, "NUMPAGES")
        for run in fp.runs:
            run.font.size = Pt(8)

    text_width = int(doc.sections[0].page_width - doc.sections[0].left_margin
                     - doc.sections[0].right_margin)
    for table in doc.tables:
        _borders(table)
        _full_width(table, text_width)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _cover_break(doc)
    _tall_signature_rows(doc)
    _keep_together(doc)
    _sort_props(doc.element.body)

    doc.save(str(path))


# --------------------------------------------------------------------------- #
def build(path: Path) -> Path:
    meta, body = read_source(path)
    if "{{TRACEABILITY_MATRIX}}" in body:
        sys.path.insert(0, str(ROOT / "tools"))
        import traceability
        problems = traceability.check(traceability.parse_urs(), traceability.parse_protocol())
        if problems:
            raise SystemExit("traceability problems:\n" + "\n".join(problems))
        body = body.replace("{{TRACEABILITY_MATRIX}}", traceability.matrix_markdown())
    md = cover_markdown(meta) + body
    out = OUT / f"{meta['doc_id']} {meta['title']}.docx"
    with tempfile.NamedTemporaryFile("w", suffix=".md", delete=False, encoding="utf-8") as tmp:
        tmp.write(md)
        tmp_path = tmp.name
    page_width_cm = 25.7 if meta.get("orientation") == "landscape" else 17.0
    subprocess.run(
        ["pandoc", tmp_path, "-f", "markdown+pipe_tables+fenced_divs+raw_attribute",
         "-t", "docx", "--reference-doc", str(REFERENCE),
         "--columns=20", "-o", str(out)],
        check=True)
    Path(tmp_path).unlink()
    postprocess(out, meta)
    return out


def main(argv: list[str]) -> int:
    OUT.mkdir(exist_ok=True)
    make_reference()
    sources = sorted(SRC.glob("*.md"))
    if argv:
        sources = [s for s in sources if any(a in s.name for a in argv)]
    for src in sources:
        print("built", build(src).relative_to(ROOT))
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
