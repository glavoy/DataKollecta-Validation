#!/usr/bin/env python3
"""Generate the standard validation test survey used by the Test Protocol (DK-VAL-004).

Writes into test-assets/:

  valtest_dictionary.xlsx   a clean data dictionary exercising every feature tested
  valtest_errors.xlsx       the same dictionary with five deliberate errors (OQ-01)
  csv/sites.csv             the CSV response list used by the `site` question
  valtest_config.json       SurveyGen configuration for version 1
  valtest_config_v2.json    configuration for version 2 (OQ-05): new surveyId, same databaseName
  valtest_dictionary_v2.xlsx  version 2: adds one question (`bednet`) to the household form

Build a package with SurveyGen from its repository root, for example:

  python main.py --config <path>/valtest_config.json

The dictionary is synthetic. It contains no study or personal data.
"""
from __future__ import annotations

import csv
import json
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent

HEADERS = ["FieldName", "QuestionType", "FieldType", "QuestionText", "MaxCharacters",
           "Responses", "LowerRange", "UpperRange", "LogicCheck", "DontKnow", "Refuse",
           "Optional", "Skip", "Comments"]

CRFS_HEADERS = ["display_order", "tablename", "displayname", "primarykey", "idconfig",
                "isbase", "linkingfield", "parenttable", "incrementfield", "requireslink",
                "repeat_count_field", "auto_start_repeat", "repeat_enforce_count",
                "display_fields", "entry_condition"]


def q(field, qtype, ftype, text="", maxc="", resp="", lo="", hi="", logic="", dk="",
      refuse="", optional="", skip="", comments=""):
    return [field, qtype, ftype, text, maxc, resp, lo, hi, logic, dk, refuse, optional,
            skip, comments]


def household(v2: bool = False):
    rows = [
        q("site", "combobox", "text", "Select the study site",
          resp="source:csv\nfile:sites.csv\ndisplay:sitename\nvalue:sitecode",
          comments="OQ-37: CSV-sourced response list"),
        q("hhid", "automatic", "text", "Household ID",
          comments="OQ-12: generated from idconfig (VT + site + 3-digit counter)"),
        q("visitdate", "date", "date", "Date of visit", lo="-1y", hi="0",
          comments="OQ-08: date range, within the last year and not in the future"),
        q("respondent", "text", "text", "Name of respondent", maxc="50"),
        q("age", "text", "text_integer", "Age of respondent in years", maxc="3", lo="18",
          hi="110", comments="OQ-08: numeric range 18-110"),
        q("sex", "radio", "integer", "Sex of respondent", resp="1:Male\n2:Female"),
        q("pregnant", "radio", "integer", "Is the respondent pregnant?",
          resp="1:Yes\n2:No", dk="TRUE",
          skip="preskip: if sex = 1, skip to weight",
          comments="OQ-10: shown only when sex = 2"),
        q("weight", "text", "text_decimal", "Weight in kg (one decimal place)", maxc="5",
          lo="2", hi="250", comments="OQ-38: decimal entry"),
        q("symptoms", "checkbox", "text", "Symptoms in the last 7 days (select all)",
          resp="1:Fever\n2:Cough\n3:Diarrhoea\n4:None"),
        q("fever_days", "text", "text_integer", "For how many days has there been fever?",
          maxc="2", lo="1", hi="30",
          skip="preskip: if symptoms 'does not contain' 1, skip to consentdate",
          comments="OQ-10: shown only when Fever is ticked"),
        q("consentdate", "date", "date", "Date consent was signed", lo="-1y", hi="0",
          logic="consentdate > visitdate; 'Consent cannot be signed after the visit date'",
          comments="OQ-09: logic check"),
        q("phone", "text", "text_integer", "Contact phone number (10 digits)", maxc="=10",
          logic="unique; 'This phone number is already registered'",
          comments="OQ-11: unique check; OQ-38: fixed length"),
        q("smoker", "radio", "integer", "Does the respondent smoke?", resp="1:Yes\n2:No",
          dk="TRUE", refuse="TRUE", comments="OQ-07: Don't know / Refuse"),
        q("arrival", "text", "hourmin", "Time of arrival (hh:mm)", maxc="=5",
          comments="OQ-38: time mask"),
    ]
    if v2:
        rows.append(q("bednet", "radio", "integer", "Does the household own a bed net?",
                      resp="1:Yes\n2:No", comments="OQ-05: added in version 2"))
    rows += [
        q("nmembers", "text", "text_integer", "Number of household members to list",
          maxc="2", lo="1", hi="20", comments="OQ-36: repeat count"),
        q("notes", "text", "text", "Notes (optional)", maxc="255", optional="TRUE",
          comments="OQ-07: optional question"),
    ]
    return rows


def member():
    return [
        q("hhid", "automatic", "text", "Household ID", comments="Linked from parent"),
        q("linenum", "automatic", "integer", "Line number", comments="Child counter"),
        q("membername", "text", "text", "Name of household member", maxc="50"),
        q("memberage", "text", "text_integer", "Age of member in years", maxc="3", lo="0",
          hi="120"),
    ]


CRFS = [
    [10, "household", "Household", "hhid",
     json.dumps({"prefix": "VT", "fields": [{"name": "site", "length": 2}],
                 "incrementLength": 3}),
     1, "hhid", "", "", "", "", "", "", "hhid,respondent", ""],
    [20, "member", "Household member", "hhid,linenum", "", 0, "hhid", "household", "linenum",
     1, "nmembers", 1, 1, "membername", ""],
]


def write_workbook(path: Path, household_rows, member_rows):
    wb = Workbook()
    ws = wb.active
    ws.title = "household_dd"
    ws.append(HEADERS)
    for r in household_rows:
        ws.append(r)
    ws2 = wb.create_sheet("member_dd")
    ws2.append(HEADERS)
    for r in member_rows:
        ws2.append(r)
    crfs = wb.create_sheet("crfs")
    crfs.append(CRFS_HEADERS)
    for r in CRFS:
        crfs.append(r)
    wb.save(path)


def broken_household():
    """Five deliberate errors, each of which SurveyGen must report (OQ-01)."""
    rows = household()
    out = []
    for r in rows:
        r = list(r)
        if r[0] == "respondent":
            r[0] = "Respondent Name"          # 1. invalid field name (capitals, space)
        if r[0] == "sex":
            r[5] = ""                         # 2. radio with no responses
        if r[0] == "pregnant":
            r[12] = "preskip: if sex = 1, skip to no_such_field"  # 3. unknown skip target
        if r[0] == "consentdate":
            r[8] = "consentdate > visit_dte; 'Bad reference'"   # 4. logic check on unknown field
        if r[0] == "age":
            r[7] = ""                         # 5. lower range without upper range
        out.append(r)
    return out


def config(version: int, excel: str, out_dir: str) -> dict:
    return {
        "excelFile": str(HERE / excel),
        "csvFiles": str(HERE / "csv"),
        "outputPath": str(HERE / out_dir),
        "surveyName": f"Validation Test Survey v{version}",
        "surveyId": f"valtest_v{version}",
        "databaseName": "valtest.sqlite",
    }


def main():
    (HERE / "csv").mkdir(exist_ok=True)
    with open(HERE / "csv" / "sites.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["sitecode", "sitename"])
        w.writerows([["01", "Site One"], ["02", "Site Two"], ["03", "Site Three"]])
    write_workbook(HERE / "valtest_dictionary.xlsx", household(), member())
    write_workbook(HERE / "valtest_dictionary_v2.xlsx", household(v2=True), member())
    write_workbook(HERE / "valtest_errors.xlsx", broken_household(), member())
    for name, cfg in (
        ("valtest_config.json", config(1, "valtest_dictionary.xlsx", "build/v1")),
        ("valtest_config_v2.json", config(2, "valtest_dictionary_v2.xlsx", "build/v2")),
        ("valtest_config_errors.json", config(1, "valtest_errors.xlsx", "build/errors")),
    ):
        (HERE / name).write_text(json.dumps(cfg, indent=2) + "\n", encoding="utf-8")
    for d in ("build/v1", "build/v2", "build/errors"):
        (HERE / d).mkdir(parents=True, exist_ok=True)
    print("test assets written to", HERE)


if __name__ == "__main__":
    main()
